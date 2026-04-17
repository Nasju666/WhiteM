from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

section = doc.sections[0]
section.top_margin    = Inches(1)
section.bottom_margin = Inches(1)
section.left_margin   = Inches(1.25)
section.right_margin  = Inches(1)

def sp(para, before=0, after=6, ls=1.5):
    pf = para.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after  = Pt(after)
    pf.line_spacing = Pt(ls * 12)

def title(text, size=16):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sp(p, 0, 6)
    r = p.add_run(text); r.bold = True
    r.font.name = 'Times New Roman'; r.font.size = Pt(size)

def chapter(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sp(p, 12, 6)
    r = p.add_run(text); r.bold = True
    r.font.name = 'Times New Roman'; r.font.size = Pt(14)

def subhead(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    sp(p, 8, 4)
    r = p.add_run(text); r.bold = True
    r.font.name = 'Times New Roman'; r.font.size = Pt(12)

def body(text, indent=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    sp(p, 0, 6)
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'; r.font.size = Pt(12)

def bullet(text, level=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    sp(p, 0, 3)
    p.paragraph_format.left_indent   = Inches(0.5 + level * 0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    r = p.add_run(u'\u2022  ' + text)
    r.font.name = 'Times New Roman'; r.font.size = Pt(12)

def pb():
    doc.add_page_break()

def center(text, size=12, italic=False, color=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sp(p, 0, 4)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'; r.font.size = Pt(size)
    r.font.italic = italic
    if color:
        r.font.color.rgb = color

# ══════════════════════════════════════════════════════════════════════════════
# TITLE PAGE
# ══════════════════════════════════════════════════════════════════════════════
for _ in range(5):
    doc.add_paragraph()

title('WHITE MARKET TRADING HUB:', 16)
title('A Blockchain-Integrated Secure Marketplace Platform', 16)
doc.add_paragraph()
center('A Capstone Project Documentation')
doc.add_paragraph()
center('Presented to')
center('College of Information Technology', 12)
center('CEBU EASTERN COLLEGE INCORPORATED', 12)
doc.add_paragraph()
center('Submitted in Fulfillment of the Requirements for the Course')
center('INTEGRATIVE PROGRAMMING AND TECHNOLOGIES', 12)
doc.add_paragraph()
center('Submitted to:')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Gladymay S. Sadorra'); r.bold = True
r.font.name = 'Times New Roman'; r.font.size = Pt(12)
doc.add_paragraph()
center('Members:')
for name in ['Amanence, Junas', '[Member 2 - Last Name, First Name]', '[Member 3 - Last Name, First Name]', '[Member 4 - Last Name, First Name]']:
    center(name)
doc.add_paragraph()
center('2025 - 2026')

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 1
# ══════════════════════════════════════════════════════════════════════════════
pb()
chapter('CHAPTER 1')
chapter('INTRODUCTION')

subhead('Background of the Study')
body('The rapid expansion of social media platforms has significantly changed how Filipinos buy and sell goods online. While these platforms provide convenience and wide reach, they lack fundamental structural safeguards to protect transacting parties from fraud. Online scams involving peer-to-peer transactions have become a persistent problem, particularly affecting students and young adults who use informal social media marketplaces to purchase affordable second-hand items.', indent=True)
body('A typical scenario involves a buyer transferring payment through a bank or e-wallet to a seller encountered on social media, only for the seller to disappear after receiving the funds. Because these platforms were not designed as transactional systems, there is no escrow mechanism, no identity verification requirement, and no formal dispute process. Victims are left with no avenue for recourse, and trust in digital commerce continues to erode.', indent=True)
body('White Market Trading Hub was developed in direct response to this problem. It is a full-stack web application combining a Django backend with Ethereum blockchain-based escrow protection. The system enforces verified seller identity through a KYC (Know Your Customer) process, locks buyer payments in a smart contract until delivery is confirmed, and provides a structured dispute resolution system administered by a platform arbiter. These mechanisms together create a trustworthy environment for online buying and selling.', indent=True)

subhead('Problem Statement')
body('Online trading through informal social media channels lacks the structural protections necessary for safe and accountable transactions. Specifically:', indent=True)
bullet('There is no escrow system — payment is transferred directly to the seller with no protection for the buyer in case of non-delivery or fraud.')
bullet('There is no verified seller identity — anonymous or fake accounts can pose as legitimate sellers without any accountability.')
bullet('There is no structured dispute resolution — victims have no platform-level mechanism to seek redress or recover funds.')
bullet('There is no transaction accountability — no immutable record links both parties to a specific transaction, making it easy for bad actors to deny involvement.')
body('These conditions lead to declining trust in online commerce and discourage participation, particularly from first-time buyers and small-scale sellers who cannot afford financial losses.', indent=True)

subhead('Objectives of the Study')
body('The primary objective of this study is to design and develop a secure, blockchain-integrated marketplace platform that addresses the identified shortcomings of informal online trading. Specifically, this study aims to:', indent=True)
bullet('Implement a KYC-based seller verification system that ensures only identity-verified users can post product listings.')
bullet('Develop a MetaMask wallet-based authentication mechanism that replaces passwords with cryptographic signature verification.')
bullet('Integrate an Ethereum smart contract escrow system that holds buyer funds securely until delivery is confirmed or a dispute is resolved.')
bullet('Build a structured dispute management system with evidence submission by both parties and administrator arbitration.')
bullet('Provide an administrative dashboard for real-time monitoring and management of users, KYC submissions, transactions, and disputes.')

subhead('Scope and Limitations')
body('White Market Trading Hub covers the full transaction lifecycle from product listing creation through payment, delivery, and dispute resolution. The platform supports three user roles: buyer, seller (KYC-verified buyer), and administrator. Blockchain integration is implemented on a local Hardhat Ethereum network during development and testing, with provisions for deployment on the Ethereum Sepolia testnet or mainnet.', indent=True)
body('The system displays prices in Philippine Peso (PHP) with ETH equivalents calculated using a fixed conversion estimate. A live cryptocurrency exchange rate API is not integrated in the current version. The platform does not include an in-built messaging or chat system; sellers may optionally provide a Facebook profile URL as an external contact channel. Partial dispute settlements (splitting escrow funds) are not supported in the current version.', indent=True)

subhead('Significance of the Study')
body('This study contributes meaningful value to the following stakeholders:', indent=True)
bullet('Buyers benefit from escrow-protected payments. Funds are locked in a smart contract and are only released upon confirmed delivery or through arbiter resolution, eliminating the financial risk of sending payment to an unknown seller.')
bullet('Sellers benefit from a trustworthy platform identity through KYC verification, which builds buyer confidence. The dispute system also provides a fair and documented process for contesting false claims.')
bullet('Administrators gain complete oversight of the platform through a centralized dashboard covering user verification, transaction monitoring, and dispute management.')
bullet('Students and researchers benefit from a practical demonstration of integrating blockchain technology with a conventional web application stack, providing a working reference for similar systems.')

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 2
# ══════════════════════════════════════════════════════════════════════════════
pb()
chapter('CHAPTER 2')
chapter('SYSTEM FEATURES AND FUNCTIONS')

subhead('System Functions')
body('System functions are the core operations that the platform must perform to fulfill its purpose. Removing any of the following functions would prevent the system from operating correctly.', indent=True)

subhead('Function 1: Wallet-Based User Authentication')
body('Users create accounts and log in using a MetaMask cryptocurrency wallet through a cryptographic nonce-signature challenge. The backend generates a unique random nonce for each login attempt. The user signs this nonce with their private key via MetaMask. The backend verifies the resulting signature against the stored wallet address and issues a JSON Web Token (JWT) upon successful verification. No passwords are stored. Each wallet address serves as a unique user identifier.', indent=True)

subhead('Function 2: KYC Seller Verification')
body('Before posting any listing, a user must complete identity verification. The three-step KYC wizard collects a government-issued ID photo (PhilSys National ID, Driver\'s License, Passport, SSS ID, UMID, or PRC ID), personal information (full name, email address, phone number, date of birth), and a selfie photo for face matching. An administrator reviews the submitted documents and either approves or rejects the submission, optionally providing a rejection reason. Approved users receive a verified seller status that is displayed as a badge on their listings and profile.', indent=True)

subhead('Function 3: Product Listing Management (CRUD)')
body('Verified sellers can create, view, update, and manage product listings. Each listing includes a title, detailed description, price in Philippine Peso, stock quantity, category, location, shipping information, optional Facebook contact URL, and one or more product photos. Listings are displayed on the marketplace to all visitors and can be searched by keyword and filtered by category. Sellers can mark listings as inactive or delete them from their dashboard.', indent=True)

subhead('Function 4: Order and Transaction Processing with Blockchain Escrow')
body('Buyers initiate a purchase by selecting a product and quantity. The system calculates the ETH equivalent of the PHP price and presents the buyer with a MetaMask transaction prompt to send that amount to the Escrow smart contract. Upon transaction confirmation, the smart contract creates an escrow record locking the funds. The blockchain transaction hash is recorded on the order. The order progresses through defined status transitions: Pending, Paid, Shipped, Delivered, and Completed. When the buyer confirms receipt, a second MetaMask transaction calls the confirmDelivery function on the contract, automatically releasing funds to the seller.', indent=True)

subhead('Function 5: Dispute Management System')
body('Either the buyer or the seller may raise a dispute on an active order in the Paid, Shipped, or Delivered state. The initiating party selects a dispute reason, provides a written description, and may upload supporting evidence files (images or PDFs). The opposing party can submit counter-evidence through the same interface. An administrator reviews all submitted evidence, adds notes, and resolves the dispute by selecting either Refund Buyer or Pay Seller. The system\'s arbiter account then automatically calls the resolveDispute function on the smart contract, transferring the locked escrow funds to the appropriate party on-chain.', indent=True)

subhead('Function 6: Admin Dashboard and Platform Management')
body('Administrators access a dedicated dashboard showing platform-wide statistics: total users, verified sellers, active listings, total orders, total revenue, pending disputes, and pending KYC submissions. Alert banners notify the administrator of items requiring immediate attention. The admin can navigate directly to the KYC review page or dispute management page from the dashboard.', indent=True)

subhead('System Features')
body('System features are enhancements that improve usability and user experience. The system continues to function if a feature is removed, but the overall quality of the user experience is reduced.', indent=True)

features = [
    ('Dark Mode / Light Mode Toggle', 'Users can switch between a dark-themed and a light-themed interface using a toggle button in the navigation bar. The selected theme is persisted across sessions using localStorage.'),
    ('Seller Ratings and Reviews', 'After an order is completed, the buyer may leave a star rating (1 to 5 stars) and written comment reviewing the seller. Average ratings are displayed on product listing cards and the seller\'s product detail page, helping future buyers make informed decisions.'),
    ('Multi-Photo Product Listings', 'Sellers can upload multiple images per listing using a drag-and-drop interface. Buyers view all photos through a gallery with clickable thumbnail navigation on the product detail page.'),
    ('Status Filter Chips', 'Orders and KYC submissions can be filtered by status using chip-style buttons, allowing users to quickly find relevant items without scrolling through an unfiltered list.'),
    ('Drag-and-Drop Evidence Upload', 'Both parties in a dispute can upload evidence files by dragging them directly onto an upload zone, providing a more accessible experience than a standard file picker input.'),
    ('Inline Admin Decision Panels', 'Administrators can reject a KYC submission or resolve a dispute using inline panels on the review cards. Rejection reasons and admin notes are entered in text areas directly on the page, eliminating the need for disruptive browser prompt dialogs.'),
    ('View Listing from Dispute', 'On the admin dispute management page, each dispute card includes a View Listing button that opens the associated product page in a new tab, allowing the administrator to inspect the product description and photos as context for their decision.'),
]
for feat_title, feat_desc in features:
    bullet(f'{feat_title} — {feat_desc}')

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 3
# ══════════════════════════════════════════════════════════════════════════════
pb()
chapter('CHAPTER 3')
chapter('SYSTEM ARCHITECTURE')

subhead('Overview')
body('White Market Trading Hub follows a three-tier architecture: a frontend presentation layer, a backend application layer, and a dual-storage layer comprising a relational database and an Ethereum blockchain. The frontend communicates with the backend through a RESTful HTTP API. The backend communicates with the blockchain through the Web3.py library. The user\'s MetaMask wallet communicates directly with the blockchain network through the browser\'s injected Ethereum provider for buyer and seller transaction signing.', indent=True)

subhead('Frontend Layer')
body('The frontend is a Single-Page Application (SPA) built with vanilla HTML, CSS, and JavaScript (ES Modules). A custom hash-based router maps URL fragments such as #/marketplace and #/orders/5 to page-rendering functions without requiring full page reloads. The interface uses CSS custom properties for theming, enabling seamless dark and light mode switching. Lucide is used for iconography throughout the application.', indent=True)
body('Key frontend modules:', indent=True)
bullet('router.js — Hash-based SPA router implementing navigate() for programmatic routing and refresh() to re-render the current page.')
bullet('pages.js — All page rendering functions including marketplace, product detail, order management, KYC wizard, seller dashboard, and all admin panels.')
bullet('api.js — Centralized HTTP client that attaches JWT tokens to every request, handles JSON serialization, and normalizes API errors.')
bullet('auth.js — Wallet connection management, nonce-based signature authentication flow, and user session persistence.')
bullet('escrow.js — MetaMask transaction helpers for the createEscrow, confirmDelivery, and disputeEscrow smart contract interactions.')

subhead('Backend Layer')
body('The backend is built with Django 4.x and Django REST Framework (DRF). It exposes a RESTful JSON API served under the /api/ prefix. Authentication uses JSON Web Tokens issued after successful MetaMask signature verification. A custom user model replaces the standard username/password model with a wallet_address primary identifier and nonce field.', indent=True)
body('Key backend components:', indent=True)
bullet('models.py — Defines all database models: User, Product, ProductImage, Order, Dispute, DisputeEvidence, KYCDocument, and SellerReview.')
bullet('serializers.py — DRF serializers for each model, handling validation, nested representations (e.g., product images within product listings), and read-only computed fields.')
bullet('views.py — API views for all endpoints organized by function: marketplace, orders, KYC, disputes, seller operations, and admin operations.')
bullet('web3_service.py — Python module using Web3.py that builds, signs with the arbiter private key, and broadcasts dispute resolution and escrow transactions to the Ethereum network.')
bullet('urls.py — URL routing configuration mapping all API endpoint paths to their corresponding views.')

subhead('Database Layer')
body('PostgreSQL is used as the primary relational database, accessed through Django\'s ORM. All application data — user accounts, product listings, orders, KYC documents, disputes, and reviews — is stored in the database. Uploaded files (product images, KYC documents, dispute evidence) are written to the server filesystem and served through Django\'s configured MEDIA_URL.', indent=True)

subhead('Blockchain Layer')
body('The Ethereum Escrow smart contract (Escrow.sol) is written in Solidity 0.8.24 and compiled and deployed using the Hardhat development framework. During development, a local Hardhat node simulates the Ethereum network. The contract maintains an on-chain mapping of escrow records, each identified by a sequential integer ID and storing the buyer address, seller address, locked ETH amount, and current status.', indent=True)
body('The smart contract defines four status states: AwaitingDelivery (funds locked, awaiting buyer confirmation), Complete (funds released to seller), Disputed (funds frozen pending arbiter review), and Resolved (arbiter has transferred funds to buyer or seller). The contract deployer (arbiter) is the only account authorized to call resolveDispute. Buyer and seller actions are signed by their MetaMask wallets and sent directly from the browser.', indent=True)

subhead('Technology Stack Summary')
tech = [
    ('Frontend', 'HTML5, CSS3 with custom properties, Vanilla JavaScript ES2022, Lucide Icons'),
    ('Backend', 'Python 3.11, Django 4.x, Django REST Framework, djangorestframework-simplejwt'),
    ('Database', 'PostgreSQL 15'),
    ('Blockchain', 'Solidity 0.8.24, Hardhat, Web3.py (backend), Ethers.js / window.ethereum (frontend)'),
    ('Wallet', 'MetaMask browser extension (EIP-1193 Ethereum provider)'),
    ('Environment Management', 'python-dotenv for secret configuration'),
]
for comp, detail in tech:
    bullet(f'{comp}: {detail}')

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 4
# ══════════════════════════════════════════════════════════════════════════════
pb()
chapter('CHAPTER 4')
chapter('DATABASE DESIGN')

subhead('Overview')
body('The database schema consists of eight primary models that collectively represent the full lifecycle of a marketplace interaction: account creation and identity verification, product listing, order processing with blockchain escrow, dispute management, and post-transaction reviews. All models are defined in Django\'s ORM and mapped to PostgreSQL tables.', indent=True)

models = [
    ('User Model', [
        ('wallet_address', 'CharField (unique)', 'Ethereum wallet address; serves as the primary login identifier'),
        ('full_name', 'CharField', 'Legal full name, populated during KYC submission'),
        ('email_address', 'EmailField', 'Contact email address'),
        ('phone_number', 'CharField', 'Contact phone number'),
        ('date_of_birth', 'DateField', 'Date of birth from KYC'),
        ('kyc_status', 'CharField', 'Verification status: none, pending, approved, or rejected'),
        ('is_staff', 'BooleanField', 'True for administrator accounts'),
        ('nonce', 'CharField', 'Random challenge value for MetaMask signature verification'),
    ]),
    ('Product Model', [
        ('seller', 'ForeignKey(User)', 'The verified seller who owns this listing'),
        ('title', 'CharField', 'Product name displayed on listings'),
        ('description', 'TextField', 'Detailed product description'),
        ('price', 'DecimalField', 'Price in Philippine Peso'),
        ('stock', 'PositiveIntegerField', 'Number of units available'),
        ('category', 'CharField', 'Product category for filtering'),
        ('status', 'CharField', 'Listing state: active, inactive, or sold_out'),
        ('location', 'CharField', 'Seller location shown on detail page'),
        ('shipping_info', 'CharField', 'Shipping method or delivery details'),
        ('facebook_url', 'URLField', 'Optional Facebook profile for direct buyer contact'),
    ]),
    ('ProductImage Model', [
        ('product', 'ForeignKey(Product)', 'The listing this image belongs to'),
        ('image', 'ImageField', 'Uploaded photo file stored on server filesystem'),
        ('order', 'PositiveIntegerField', 'Integer for controlling display order in the gallery'),
    ]),
    ('Order Model', [
        ('buyer', 'ForeignKey(User)', 'The user who placed the order'),
        ('product', 'ForeignKey(Product)', 'The product being purchased'),
        ('quantity', 'PositiveIntegerField', 'Number of units ordered'),
        ('total_price', 'DecimalField', 'Total PHP price locked at the time of order'),
        ('status', 'CharField', 'Order lifecycle state: pending, paid, shipped, delivered, completed, disputed, or cancelled'),
        ('escrow_id', 'IntegerField (nullable)', 'The on-chain escrow record ID returned by the smart contract'),
        ('tx_hash', 'CharField (nullable)', 'Blockchain transaction hash of the escrow creation transaction'),
    ]),
    ('Dispute Model', [
        ('order', 'OneToOneField(Order)', 'The order that is under dispute'),
        ('opened_by', 'ForeignKey(User)', 'The buyer or seller who raised the dispute'),
        ('reason', 'CharField', 'Reason code: not_received, not_as_described, damaged, unauthorized, or other'),
        ('description', 'TextField', 'Detailed written explanation of the issue'),
        ('resolution', 'CharField', 'Outcome: pending, refund_buyer, or pay_seller'),
        ('admin_notes', 'TextField', 'Notes recorded by the administrator during review'),
        ('resolved_by', 'ForeignKey(User, nullable)', 'The administrator who closed the dispute'),
        ('resolved_at', 'DateTimeField (nullable)', 'Timestamp when the dispute was resolved'),
    ]),
    ('DisputeEvidence Model', [
        ('dispute', 'ForeignKey(Dispute)', 'The dispute this evidence is associated with'),
        ('uploaded_by', 'ForeignKey(User)', 'The party (buyer or seller) who submitted this file'),
        ('file', 'FileField', 'Uploaded evidence file: image or PDF'),
        ('description', 'TextField', 'Description of what the evidence shows'),
        ('uploaded_at', 'DateTimeField', 'Timestamp of upload'),
    ]),
    ('KYCDocument Model', [
        ('user', 'ForeignKey(User)', 'The user submitting for verification'),
        ('document_type', 'CharField', 'Type: philsys_id, drivers_license, passport, sss_id, umid, prc_id, or selfie'),
        ('file', 'ImageField', 'Uploaded document image file'),
        ('status', 'CharField', 'Review state: pending, approved, or rejected'),
        ('notes', 'TextField', 'Rejection reason or admin review notes'),
        ('uploaded_at', 'DateTimeField', 'Timestamp of document submission'),
        ('reviewed_at', 'DateTimeField (nullable)', 'Timestamp of admin review'),
    ]),
    ('SellerReview Model', [
        ('order', 'OneToOneField(Order)', 'The completed order this review is linked to — enforces one review per order'),
        ('reviewer', 'ForeignKey(User)', 'The buyer writing the review'),
        ('seller', 'ForeignKey(User)', 'The seller being reviewed'),
        ('rating', 'IntegerField', 'Star rating from 1 to 5'),
        ('comment', 'TextField', 'Written review text'),
        ('created_at', 'DateTimeField', 'Timestamp when the review was submitted'),
    ]),
]

for model_name, fields in models:
    subhead(model_name)
    for fname, ftype, fdesc in fields:
        bullet(f'{fname} ({ftype}) — {fdesc}')

subhead('Key Relationships')
bullet('A User (as seller) has many Products; each Product belongs to exactly one seller.')
bullet('A Product has many ProductImages ordered by an integer field for gallery display.')
bullet('An Order links one buyer User to one Product and stores the escrow reference data from the blockchain.')
bullet('A Dispute is a one-to-one extension of an Order, created only when the transaction is contested.')
bullet('DisputeEvidence records belong to a Dispute and are associated with the uploading User, allowing buyer and seller submissions to be displayed separately.')
bullet('KYCDocument records belong to a User; multiple documents (ID photo and selfie) are stored as separate records with different document_type values.')
bullet('A SellerReview is a one-to-one extension of a completed Order, enforcing the rule that each transaction can receive only one review.')

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 5
# ══════════════════════════════════════════════════════════════════════════════
pb()
chapter('CHAPTER 5')
chapter('SCREENSHOTS OF THE SYSTEM')

screens = [
    ('Landing Page', 'The landing page presents the platform with a hero section communicating the core value proposition of secure, blockchain-backed trading. Navigation links for Marketplace, Dashboard, Orders, and KYC are displayed. The page includes a Connect Wallet button that initiates the MetaMask authentication flow.'),
    ('Marketplace', 'The marketplace displays all active product listings in a responsive card grid. Each card shows the primary product image, product title, PHP price with ETH equivalent, seller verification badge, category, and average star rating. A search bar and category filter chips allow users to narrow listings.'),
    ('Product Detail Page', 'The product detail page shows a full image gallery with thumbnail navigation, complete description, location, shipping information, and the seller\'s verification status and average review rating. A Buy Now section with quantity selection is displayed for buyers. Customer reviews and ratings from previous buyers are shown at the bottom.'),
    ('KYC Verification Wizard - Step 1', 'Step 1 of the KYC wizard shows a drag-and-drop zone for uploading a government-issued ID, a document type dropdown, and an informational tip about photo quality requirements. Completed steps are marked with a green checkmark in the step tracker.'),
    ('KYC Verification Wizard - Step 2', 'Step 2 collects the user\'s full name, email address, phone number, and date of birth in a two-column form layout. Field values are retained when navigating between wizard steps.'),
    ('KYC Verification Wizard - Step 3', 'Step 3 displays a drag-and-drop selfie upload zone and a Submission Summary table showing all previously entered information for final review before submission.'),
    ('Create Listing Page', 'The create listing page is organized into four sections: Basic Info (title, description, price, stock, category), Media (drag-and-drop multi-image upload), Shipping and Location (location field, shipping details), and Contact (optional Facebook URL). A live preview of the primary photo is shown.'),
    ('Orders Page', 'The orders page lists all user orders with status filter chips (All, Pending, Paid, Shipped, Delivered, Completed, Disputed). Each order card displays a color-coded status icon, product name, quantity, order date, PHP price with ETH equivalent, and a status pill badge. Clicking a card navigates to the order detail page.'),
    ('Order Detail Page', 'The order detail page shows full order information and action buttons appropriate to the user\'s role and the current order status. Below the details, a dispute section appears showing the active dispute, evidence from both buyer and seller in separated panels, and admin resolution notes when applicable.'),
    ('Dispute Submission Form', 'The dispute form is embedded inline on the order detail page. It presents reason options as visual radio cards with icons, a text area for the description, and a drag-and-drop evidence upload zone with a cancel button.'),
    ('Seller Dashboard', 'The seller dashboard shows four stat cards (total listings, active listings, total sales, revenue) and a filterable list of the seller\'s product listings and orders. Each listing card includes the product thumbnail, title, price, stock level, status, and links to edit or view the listing.'),
    ('Admin Dashboard', 'The admin dashboard shows seven stat cards with color-coded icons. Alert banners highlight pending disputes and KYC submissions with direct action links. An orders-by-status horizontal bar chart shows the distribution of order states as percentages.'),
    ('Admin KYC Review Page', 'The KYC review page shows stats pills (Pending, Approved, Rejected counts) and lists each submission as a card. Pending cards display the wallet address, submission time, personal information in a two-column grid, document links, and Approve/Reject buttons. Clicking Reject reveals an inline rejection reason text area.'),
    ('Admin Dispute Management Page', 'Each dispute card shows the product title, order number, buyer and seller wallets, dispute reason, description, and evidence organized into separate Buyer and Seller columns. A View Listing button opens the product page. The resolve panel contains an admin notes textarea and Refund Buyer / Pay Seller buttons.'),
]

for i, (fig_title, desc) in enumerate(screens, 1):
    subhead(f'Figure {i}: {fig_title}')
    body(desc, indent=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sp(p, 4, 12)
    r = p.add_run(f'[Insert screenshot of {fig_title} here]')
    r.font.name = 'Times New Roman'; r.font.size = Pt(11)
    r.font.italic = True
    r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

# ══════════════════════════════════════════════════════════════════════════════
# CHAPTER 6
# ══════════════════════════════════════════════════════════════════════════════
pb()
chapter('CHAPTER 6')
chapter('USER MANUAL')

subhead('System Requirements')
body('The following are required to access and use White Market Trading Hub:', indent=True)
bullet('A modern web browser — Google Chrome or Brave Browser is recommended for full MetaMask compatibility.')
bullet('MetaMask browser extension — installed, set up, and configured with at least one Ethereum wallet account.')
bullet('Sufficient ETH balance in the connected wallet to cover product payment amounts and gas fees for blockchain transactions.')
bullet('An active internet connection, or access to the local development server if running in a testing environment.')

subhead('Part 1: Getting Started — Connecting Your Wallet')
for step in [
    'Open the White Market Trading Hub URL in your browser.',
    'Click the Connect Wallet button in the top navigation bar.',
    'MetaMask will open and ask you to select an account. Choose your wallet and click Connect.',
    'MetaMask will prompt you to sign a message to verify your wallet ownership. Click Sign. No ETH is spent for this step.',
    'You are now logged in. Your shortened wallet address will appear in the top-right corner of the navigation bar.',
]:
    bullet(step)

subhead('Part 2: Completing KYC Verification (Required to Sell)')
body('Before you can post any listings, you must complete identity verification.', indent=True)
for step in [
    'Click the KYC button in the navigation bar.',
    'Step 1 — ID Document: Choose your document type from the dropdown. Click or drag a clear, well-lit photo of your ID into the upload zone. All four corners of the ID must be visible. Click Next.',
    'Step 2 — Personal Information: Enter your full name, email address, phone number, and date of birth exactly as they appear on the submitted ID. Click Next.',
    'Step 3 — Face Verification: Upload a clear, front-facing selfie photo. Review the Submission Summary to confirm all details are correct. Click Submit for Verification.',
    'Your submission will be placed under review. The KYC page will display an Under Review status. Once an administrator approves your submission, your status changes to Approved and you may begin posting listings.',
    'If rejected, the page will show the rejection reason. Click Resubmit Verification to go through the wizard again with corrected documents.',
]:
    bullet(step)

subhead('Part 3: Creating a Product Listing')
for step in [
    'Click Sell in the navigation bar. You must be KYC-approved to access this page.',
    'Basic Info: Enter the product title, description, price in PHP, available quantity, and select a category.',
    'Media: Upload one or more product photos by clicking or dragging image files into the upload zone.',
    'Shipping and Location: Enter your location and describe your shipping method (e.g., Grab Express, LBC, meetup only).',
    'Contact (optional): Paste your Facebook profile URL if you wish to provide an external contact channel for buyers.',
    'Click Publish Listing. Your listing will immediately appear on the marketplace.',
]:
    bullet(step)

subhead('Part 4: Placing an Order (Buyer)')
for step in [
    'Browse the Marketplace and click on a product to view its detail page.',
    'Select the desired quantity and click Buy Now.',
    'MetaMask will open showing a transaction sending ETH (equivalent to the PHP price) to the escrow smart contract. Review the amount carefully.',
    'Click Confirm in MetaMask to send the transaction. Wait for blockchain confirmation — this may take a few seconds.',
    'Your order is created with status Paid. The seller will receive a notification.',
    'When the seller marks the order as Shipped, you will see the status update on your Orders page.',
    'Once you receive the item, open the order and click Confirm Delivery. Sign the MetaMask transaction. The escrow funds will be automatically released to the seller and the order will be marked Completed.',
]:
    bullet(step)

subhead('Part 5: Managing an Order (Seller)')
for step in [
    'Navigate to your Dashboard or click Orders in the navigation bar.',
    'Find a Paid order and click on it to view the details including the buyer information and delivery address.',
    'Once you have dispatched the item, click Mark as Shipped.',
    'After the buyer confirms receipt, the order status will change to Completed and the escrowed ETH will be transferred to your wallet automatically.',
]:
    bullet(step)

subhead('Part 6: Opening a Dispute')
body('If there is a problem with an order — such as non-delivery or an item that does not match the description — either the buyer or seller may raise a dispute.', indent=True)
for step in [
    'Go to Orders and click on the affected order.',
    'Click Open Dispute.',
    'Select the reason that best describes the issue from the reason cards (Item Not Received, Not As Described, Damaged, Unauthorized Transaction, or Other).',
    'Write a detailed description of the problem in the text area.',
    'Optionally upload supporting evidence (photos or PDF documents) using the evidence upload zone.',
    'Click Submit Dispute. The order status changes to Disputed and both parties can see the active dispute on the order page.',
    'The opposing party may navigate to the same order and submit counter-evidence in their own evidence section.',
    'An administrator will review all evidence and resolve the dispute. You will see the resolution and admin notes on the order page once decided.',
]:
    bullet(step)

subhead('Part 7: Administrator — Reviewing KYC Submissions')
for step in [
    'Log in with an administrator account. Click KYC Review in the navigation bar.',
    'The stats pills at the top show the count of Pending, Approved, and Rejected submissions.',
    'Click the document links on each Pending card to open the uploaded ID photos and selfie in a new tab.',
    'Verify that the document is valid, legible, and matches the submitted personal information.',
    'Click Approve if the documents are satisfactory. The user\'s KYC status will be updated and they will be able to create listings.',
    'Click Reject if the documents are invalid or do not match. An inline text area will appear — enter a clear rejection reason and click Confirm Reject. The user will see this reason when they return to the KYC page.',
]:
    bullet(step)

subhead('Part 8: Administrator — Resolving Disputes')
for step in [
    'Click Disputes in the navigation bar.',
    'Review each Pending dispute card. The card shows both parties\' wallet addresses, the dispute reason and description, and evidence organized by Buyer and Seller.',
    'Click View Listing to inspect the associated product detail page and verify whether the seller\'s listing accurately describes what was sold.',
    'Review all uploaded evidence from both parties.',
    'Enter any administrative notes in the text area within the Admin Decision panel.',
    'Click Refund Buyer if the buyer\'s claim is valid and they should receive their funds back.',
    'Click Pay Seller if the seller has fulfilled their obligation and the dispute is not justified.',
    'The system will automatically execute the on-chain escrow resolution using the platform\'s arbiter account, transferring the locked funds to the appropriate party.',
]:
    bullet(step)

doc.save('White_Market_Manuscript.docx')
print('Done.')
